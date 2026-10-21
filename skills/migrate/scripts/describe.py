# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl>=3.1", "oletools>=0.60"]
# ///
"""Describes an Excel workbook for documenting its process: sheets with their size, header and formulas, defined names, and the full VBA code of every module.

Usage: uv run describe.py <workbook> [--max-formulas N]. Prints Markdown; nothing is changed.
"""

import argparse
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def sheets(path: Path, max_formulas: int) -> list[str]:
    wb = load_workbook(path, data_only=False)
    out = ["## Hojas", ""]
    for ws in wb.worksheets:
        header = [c.value for c in ws[1]] if ws.max_row else []
        formulas = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        out.append(f"### {ws.title}")
        out.append(f"- Tamaño: {ws.max_row} filas × {ws.max_column} columnas ({ws.dimensions})")
        out.append(f"- Cabecera (fila 1): {header}")
        if formulas:
            shapes = Counter(re.sub(r"\d+", "#", value) for _, value in formulas)  # B2/1000 and B3/1000 are the same formula filled down
            out.append(f"- Fórmulas: {len(formulas)}; patrones más frecuentes: {shapes.most_common(3)}")
            out += [f"  - `{coord}`: `{value}`" for coord, value in formulas[:max_formulas]]
        out.append("")
    if wb.defined_names:
        out += ["## Nombres definidos", ""] + [f"- `{name}`: `{dn.attr_text}`" for name, dn in wb.defined_names.items()] + [""]
    return out


def vba(path: Path) -> list[str]:
    from oletools.olevba import VBA_Parser

    parser = VBA_Parser(str(path))
    try:
        if not parser.detect_vba_macros():
            return ["## VBA", "", "El libro no contiene macros.", ""]
        out = ["## VBA", ""]
        for _, _, name, code in parser.extract_macros():
            body = "\n".join(line for line in code.splitlines() if not line.startswith("Attribute "))
            if body.strip():
                out += [f"### {name}", "", "```vba", body.strip(), "```", ""]
        return out
    finally:
        parser.close()


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("workbook", type=Path)
    p.add_argument("--max-formulas", type=int, default=5)
    a = p.parse_args()
    print("\n".join([f"# {a.workbook.name}", ""] + sheets(a.workbook, a.max_formulas) + vba(a.workbook)))


if __name__ == "__main__":
    main()
