"""Settlement of one installation, ported from the Settle macro (macro/Settlement.bas).

Rules kept exactly as the macro applies them, so the characterization tests can pin them:
- billable energy per hour: measured kWh / 1000 × (1 + losses);
- amount per hour: round(MWh × (price + toll), 2) before summing the day;
- an hour without a measure takes the average of the day's measured hours;
- a day without any measure settles to zero and is flagged;
- an hour without a price is settled at price 0 (the macro does this silently; `missing_prices` makes it visible).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class Inputs:
    losses: float
    toll: float
    measures: list[tuple[date, int, float | None]]  # (day, hour, kWh or None), in sheet order
    prices: dict[tuple[date, int], float]


@dataclass(frozen=True)
class DayResult:
    day: date
    hours: int
    estimated_hours: int
    energy_mwh: float
    amount_eur: float
    without_measures: bool
    missing_prices: list[int] = field(default_factory=list)


def _as_date(value: date | datetime) -> date:
    return value.date() if isinstance(value, datetime) else value


def read_inputs(path: Path) -> Inputs:
    """Reads the Parametros, Medidas and Precios sheets. Column C holds the raw kWh; column D is a formula and is ignored on purpose."""
    wb = load_workbook(path, data_only=True, read_only=True)
    params = wb["Parametros"]
    losses = float(params["B2"].value)
    toll = float(params["B3"].value)
    measures = [(_as_date(d), int(h), None if kwh is None else float(kwh)) for d, h, kwh, *_ in wb["Medidas"].iter_rows(min_row=2, values_only=True) if d is not None]
    prices = {(_as_date(d), int(h)): float(p) for d, h, p in wb["Precios"].iter_rows(min_row=2, max_col=3, values_only=True) if d is not None}
    return Inputs(losses, toll, measures, prices)


def settle_day(day: date, rows: list[tuple[int, float | None]], inputs: Inputs) -> DayResult:
    measured = [kwh for _, kwh in rows if kwh is not None]
    average = sum(measured) / len(measured) if measured else 0.0
    energy = amount = 0.0
    missing_prices = []
    for hour, kwh in rows:
        mwh = (average if kwh is None else kwh) / 1000 * (1 + inputs.losses)
        energy += mwh
        price = inputs.prices.get((day, hour))
        if price is None:
            missing_prices.append(hour)
            price = 0.0
        amount += round(mwh * (price + inputs.toll), 2)
    return DayResult(day, len(rows), len(rows) - len(measured), energy, amount, not measured, missing_prices)


def settle(inputs: Inputs) -> list[DayResult]:
    """One result per day, in the order the days appear in Medidas."""
    days: dict[date, list[tuple[int, float | None]]] = {}
    for day, hour, kwh in inputs.measures:
        days.setdefault(day, []).append((hour, kwh))
    return [settle_day(day, rows, inputs) for day, rows in days.items()]
